"""Excel/CSV student import."""

from __future__ import annotations

import csv
import io
import re
import uuid
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import AppError
from app.core.security import normalize_email
from app.models.campus import Campus, CampusSettings
from app.models.college import CollegeDomain
from app.models.enums import (
    DomainStatus,
    ImportBatchStatus,
    ImportRowStatus,
    MembershipRole,
    StudentStatus,
)
from app.models.student import Student, StudentImportBatch, StudentImportRow
from app.services.auth import ensure_membership, get_or_create_user

_HEADER_ALIASES = {
    "name": {"name", "full_name", "student_name"},
    "srn": {"srn", "usn", "roll", "roll_no", "roll_number"},
    "email": {"email", "academic_email", "student_email"},
    "lat": {"lat", "latitude", "home_lat"},
    "lng": {"lng", "lon", "long", "longitude", "home_lng", "home_lon"},
}


def _norm_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.strip().lower()).strip("_")


def _map_headers(headers: list[str]) -> dict[str, int]:
    normalized = [_norm_header(h) for h in headers]
    mapping: dict[str, int] = {}
    for field, aliases in _HEADER_ALIASES.items():
        for idx, header in enumerate(normalized):
            if header in aliases:
                mapping[field] = idx
                break
    missing = [f for f in ("name", "srn", "email", "lat", "lng") if f not in mapping]
    if missing:
        raise AppError(
            "VALIDATION_ERROR",
            f"Missing required columns: {', '.join(missing)}.",
            status_code=422,
        )
    return mapping


def _parse_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    elif lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = [[cell if cell is not None else "" for cell in row] for row in ws.iter_rows(values_only=True)]
    else:
        raise AppError(
            "VALIDATION_ERROR",
            "Unsupported file type. Upload .csv or .xlsx.",
            status_code=422,
        )

    if not rows:
        raise AppError("VALIDATION_ERROR", "File is empty.", status_code=422)

    headers = [str(h) for h in rows[0]]
    mapping = _map_headers(headers)
    parsed: list[dict[str, Any]] = []
    for i, row in enumerate(rows[1:], start=2):
        if not row or all(str(c).strip() == "" for c in row):
            continue
        def cell(field: str) -> str:
            idx = mapping[field]
            if idx >= len(row):
                return ""
            return str(row[idx]).strip()

        parsed.append(
            {
                "row_number": i,
                "name": cell("name"),
                "srn": cell("srn"),
                "email": cell("email"),
                "lat": cell("lat"),
                "lng": cell("lng"),
            }
        )
    if not parsed:
        raise AppError("VALIDATION_ERROR", "No data rows found.", status_code=422)
    return parsed


def _active_domains(db: Session, college_id: uuid.UUID) -> set[str]:
    rows = db.scalars(
        select(CollegeDomain).where(
            CollegeDomain.college_id == college_id,
            CollegeDomain.status == DomainStatus.ACTIVE,
        )
    ).all()
    return {r.domain.lower() for r in rows}


def import_students_file(
    db: Session,
    *,
    college_id: uuid.UUID,
    campus_id: uuid.UUID,
    uploaded_by_user_id: uuid.UUID,
    filename: str,
    content: bytes,
) -> StudentImportBatch:
    campus = db.scalar(
        select(Campus).where(Campus.id == campus_id, Campus.college_id == college_id)
    )
    if campus is None:
        raise AppError("NOT_FOUND", "Campus not found.", status_code=404)

    settings = db.scalar(
        select(CampusSettings).where(
            CampusSettings.college_id == college_id,
            CampusSettings.campus_id == campus_id,
        )
    )
    if settings is None:
        raise AppError(
            "CONFLICT",
            "Campus settings required before student import.",
            status_code=409,
        )

    domains = _active_domains(db, college_id)
    if not domains:
        raise AppError(
            "CONFLICT",
            "At least one active college domain is required before import.",
            status_code=409,
        )

    parsed = _parse_rows(filename, content)
    batch = StudentImportBatch(
        college_id=college_id,
        campus_id=campus_id,
        uploaded_by_user_id=uploaded_by_user_id,
        filename=filename,
        status=ImportBatchStatus.PROCESSING,
        row_count=len(parsed),
    )
    db.add(batch)
    db.flush()

    valid = 0
    errors = 0
    for raw in parsed:
        row_status = ImportRowStatus.VALID
        error_message = None
        student_id = None
        try:
            name = raw["name"]
            srn = raw["srn"].strip().upper()
            email = normalize_email(raw["email"])
            lat = float(raw["lat"])
            lng = float(raw["lng"])
            if not name or not srn or not email:
                raise ValueError("name, srn, and email are required")
            if not (-90 <= lat <= 90 and -180 <= lng <= 180):
                raise ValueError("coordinates out of range")
            domain = email.split("@", 1)[1]
            if domain not in domains:
                raise ValueError(f"email domain '{domain}' is not an approved college domain")

            existing = db.scalar(
                select(Student).where(
                    Student.college_id == college_id,
                    (Student.srn == srn) | (Student.email == email),
                )
            )
            if existing is not None:
                raise ValueError("student with same SRN or email already exists")

            user = get_or_create_user(db, email, name)
            ensure_membership(
                db,
                college_id=college_id,
                user_id=user.id,
                role=MembershipRole.STUDENT,
            )
            student = Student(
                college_id=college_id,
                campus_id=campus_id,
                user_id=user.id,
                import_batch_id=batch.id,
                srn=srn,
                email=email,
                full_name=name,
                home_lat=lat,
                home_lng=lng,
                status=StudentStatus.IMPORTED,
            )
            db.add(student)
            db.flush()
            student_id = student.id
            valid += 1
        except Exception as exc:  # noqa: BLE001 — row-level capture
            row_status = ImportRowStatus.INVALID
            error_message = str(exc)
            errors += 1

        db.add(
            StudentImportRow(
                batch_id=batch.id,
                college_id=college_id,
                row_number=int(raw["row_number"]),
                raw_payload=raw,
                status=row_status,
                error_message=error_message,
                student_id=student_id,
            )
        )

    batch.valid_count = valid
    batch.error_count = errors
    batch.status = ImportBatchStatus.COMPLETED
    db.commit()
    db.refresh(batch)
    return batch


def get_batch(
    db: Session, college_id: uuid.UUID, batch_id: uuid.UUID
) -> StudentImportBatch:
    batch = db.scalar(
        select(StudentImportBatch).where(
            StudentImportBatch.id == batch_id,
            StudentImportBatch.college_id == college_id,
        )
    )
    if batch is None:
        raise AppError("NOT_FOUND", "Import batch not found.", status_code=404)
    return batch


def list_students(
    db: Session,
    college_id: uuid.UUID,
    *,
    campus_id: uuid.UUID | None = None,
) -> list[Student]:
    stmt = select(Student).where(Student.college_id == college_id).order_by(Student.srn)
    if campus_id is not None:
        stmt = stmt.where(Student.campus_id == campus_id)
    return list(db.scalars(stmt).all())
